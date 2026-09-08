from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
OUT_PATH = OUT_DIR / "Manual_de_uso_sistema_de_nomina_Keysar_Cosmetics.docx"

INK = "1F1B18"
BROWN = "3B2B22"
COPPER = "9B6A43"
SAND = "F3E9DE"
PALE = "FAF7F3"
GRID = "D9D9D9"
GREEN = "007A5A"
RED = "B4233F"
WHITE = "FFFFFF"
MUTED = "6B625C"


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        shade_cell(cell, BROWN)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, size=9, bold=True, color=WHITE)
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        for cidx, value in enumerate(row_data):
            cell = row.cells[cidx]
            cell.text = ""
            shade_cell(cell, WHITE if ridx % 2 == 0 else PALE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.5, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.5, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=INK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=INK)
    return p


def add_step(doc, label, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    r1 = p.add_run(label + " ")
    set_run_font(r1, size=10.3, bold=True, color=INK)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.3, color=INK)
    return p


def add_status_line(doc, label, text, color):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.2, bold=True, color=color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.2, color=INK)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run("Keysar Cosmetics   |   Manual de uso   |   Página ")
    set_run_font(r, size=8, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr)
    r._r.append(fld_char2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")

    heading_sizes = {1: 19, 2: 14, 3: 11.5}
    for level, size in heading_sizes.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(65)
    r = p.add_run("KEYSAR")
    set_run_font(r, name="Georgia", size=24, bold=True, color=BROWN)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(42)
    r = p2.add_run("COSMETICS  ·  PAYROLL")
    set_run_font(r, size=10, bold=True, color=COPPER)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    title.add_run("Manual de uso del sistema de nómina")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(44)
    r = subtitle.add_run("Operación administrativa y portal del personal")
    set_run_font(r, size=14, color=MUTED)

    meta = add_table(
        doc,
        ["Documento", "Versión", "Fecha", "Alcance"],
        [["Manual de usuario", "1.0", "7 de septiembre de 2026", "Prototipo frontend"]],
        widths=[1.55, 0.8, 1.65, 1.75],
        font_size=9.5,
    )
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    add_body(
        doc,
        "Este manual explica cómo configurar, calcular, revisar, autorizar, conciliar y exportar la nómina, además de cómo usa el personal su portal privado.",
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run("Uso interno y confidencial")
    set_run_font(r, size=9, bold=True, color=COPPER)
    doc.add_page_break()


def build_manual():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    add_footer(section)

    add_cover(doc)

    add_heading(doc, "Alcance de este manual", 1)
    add_body(
        doc,
        "El manual cubre el portal administrativo para usuarios máster y el portal privado del personal. Describe la versión de demostración disponible en el proyecto de nómina de Keysar Cosmetics.",
    )
    add_body(
        doc,
        "Importante  En la versión actual los datos se conservan únicamente en la memoria de la sesión. Recargar la aplicación o usar Restaurar demostración devuelve los datos iniciales. No se realizan transferencias bancarias reales, no existe persistencia productiva y la autenticación es simulada.",
        bold_lead="Importante  ",
    )

    add_heading(doc, "Contenido", 1)
    contents = [
        "1 Acceso y seguridad",
        "2 Navegación y controles comunes",
        "3 Orden recomendado de configuración",
        "4 Gestión de empleados",
        "5 Catálogos y estructura de nómina",
        "6 Módulos de nómina",
        "7 Cálculo y cierre de una nómina",
        "8 Costo social e ISR",
        "9 Movimientos de nómina",
        "10 Préstamos y adelantos",
        "11 Recibos y autorizaciones del personal",
        "12 Reportes y conciliación",
        "13 Dispersión y exportaciones",
        "14 Procesos especiales",
        "15 Controles antes del pago",
        "16 Solución de problemas",
        "17 Glosario y rutas del sistema",
    ]
    for item in contents:
        add_bullet(doc, item)
    doc.add_page_break()

    add_heading(doc, "1 Acceso y seguridad", 1)
    add_body(
        doc,
        "El sistema separa el acceso administrativo del acceso personal. El usuario máster ve todos los módulos. Los demás usuarios sólo pueden abrir los módulos autorizados para su rol y su información personal.",
    )
    add_heading(doc, "Perfiles de acceso", 2)
    add_table(
        doc,
        ["Perfil", "Uso", "Alcance"],
        [
            ["Usuario máster", "Configura y opera la nómina", "Acceso obligatorio a todos los módulos, reportes y autorizaciones"],
            ["Empleado", "Consulta y confirma su pago", "Sólo su portal, sus recibos, solicitudes y módulos autorizados"],
            ["Gerencia", "Consulta ventas y comisión gerencial", "Recibo personal de ventas y recibo gerencial separado cuando cierra el mes"],
        ],
        widths=[1.25, 2.05, 3.45],
    )
    add_heading(doc, "Accesos de demostración", 2)
    add_table(
        doc,
        ["Perfil", "Usuario", "Clave", "Código secundario"],
        [
            ["Máster", "MASTER DEMO", "NOMINA2026", "2580"],
            ["Empleado", "VENDEDOR DEMO", "VENTAS2026", "1470"],
        ],
        widths=[1.15, 1.7, 1.7, 1.65],
    )
    add_heading(doc, "Inicio de sesión", 2)
    add_step(doc, "Seleccione el perfil", "Elija Empleado o Máster en la pantalla de acceso.")
    add_step(doc, "Capture las credenciales", "Escriba el usuario y la clave principal. El navegador no debe completar el segundo código.")
    add_step(doc, "Complete la verificación", "Ingrese el código de cuatro dígitos con el teclado visual y pulse Verificar e ingresar.")
    add_step(doc, "Confirme el destino", "El perfil máster abre el control general. Un empleado abre únicamente su portal y sus permisos autorizados.")
    add_heading(doc, "Contraseña temporal y segundo código", 2)
    add_body(
        doc,
        "Al terminar el alta de un empleado, el sistema genera una contraseña temporal. En su primer acceso, el empleado debe cambiarla y elegir su propio código secundario. Roles y accesos muestra el estado del cambio sin volver a revelar la contraseña o el código.",
    )
    add_status_line(doc, "Seguridad", "La sesión se cierra después de 3 minutos sin teclado, puntero, toque o desplazamiento. Todos los diálogos se cierran y se deben repetir los dos pasos de acceso.", COPPER)
    add_status_line(doc, "Confidencialidad", "No comparta credenciales, CLABE, recibos ni exportaciones. El acceso por URL también se bloquea cuando el rol no tiene permiso.", RED)

    add_heading(doc, "2 Navegación y controles comunes", 1)
    add_body(
        doc,
        "Los menús viven en la parte superior. Personal, Nómina, Operación, Configuración y Reportes abren sus submenús con un clic. En pantallas pequeñas, el mismo contenido aparece en un panel compacto.",
    )
    add_table(
        doc,
        ["Menú", "Submenús principales"],
        [
            ["Personal", "Empleados"],
            ["Nómina", "Consolidado, Salario fijo, Especialistas, Comisiones, Comisión de kiosco, Honorarios, Dispersión de nómina"],
            ["Operación", "Cálculo de comisiones, Movimientos, Préstamos"],
            ["Configuración", "Periodos y conceptos, Módulos de nómina, Sucursales, Puestos, Esquemas, Bonos y multas, Viáticos, Roles y accesos"],
            ["Reportes", "Gastos por puesto, Desglose por sucursal, Recibos, Recibos gerenciales"],
        ],
        widths=[1.35, 5.25],
    )
    add_heading(doc, "Controles que se repiten", 2)
    add_bullet(doc, "Búsqueda  Filtra por nombre, puesto, sucursal, rol o concepto según el módulo.")
    add_bullet(doc, "Filas  Permite mostrar 20, 40, 60 o todos los registros. Use Anterior y Siguiente para cambiar de página.")
    add_bullet(doc, "Periodo  Selecciona quincena, mes, trimestre o año cuando el reporte lo permite.")
    add_bullet(doc, "Interruptor  Activa o desactiva una regla. Revise la etiqueta del periodo antes de cambiarlo.")
    add_bullet(doc, "Casillas  Permiten seleccionar uno, varios o todos los registros visibles o filtrados.")
    add_bullet(doc, "Exportación  Los iconos de impresora, PDF y Excel usan la misma información visible y los filtros activos.")
    add_bullet(doc, "Tema  Cambia entre visualización clara y oscura sin modificar los datos.")
    add_heading(doc, "Estados visuales", 2)
    add_status_line(doc, "Verde", "Proceso correcto, aprobación completa o conciliación sin diferencias.", GREEN)
    add_status_line(doc, "Rojo", "Falta una aprobación, existe una diferencia o una regla impide continuar.", RED)
    add_status_line(doc, "Borrador", "La información puede modificarse y aún no está disponible para dispersión.", COPPER)

    doc.add_page_break()
    add_heading(doc, "3 Orden recomendado de configuración", 1)
    add_body(doc, "Use este orden al preparar una operación nueva o al iniciar un periodo de nómina.")
    config_steps = [
        ("Sucursales", "Registre o sincronice los puntos de venta y confirme cuáles están activos."),
        ("Puestos", "Cree el catálogo oficial. Los puestos se eligen de la lista y no se escriben libremente."),
        ("Módulos de nómina", "Confirme los módulos base o cree módulos personalizados con sus conceptos y puestos."),
        ("Periodos y conceptos", "Defina cortes, reglas fiscales y políticas de préstamos y adelantos."),
        ("Esquemas", "Configure rangos de comisión y asigne vigencias por empleado."),
        ("Esquemas por sucursal", "Defina comisión de kiosco, sucursales incluidas y gerencia vigente."),
        ("Bonos multas y viáticos", "Configure conceptos, importes, escalas y autorizaciones."),
        ("Roles y accesos", "Asigne permisos por módulo. El rol máster siempre conserva acceso total."),
        ("Empleados", "Registre datos laborales, bancarios, vigencia y distribución del costo."),
        ("Cálculo y revisión", "Calcule la nómina, concilie, cierre, recopile autorizaciones y genere la dispersión."),
    ]
    for label, text in config_steps:
        add_step(doc, label, text)

    add_heading(doc, "4 Gestión de empleados", 1)
    add_heading(doc, "Alta de un empleado", 2)
    add_step(doc, "Abra Personal y Empleados", "Use Nuevo empleado para iniciar el registro.")
    add_step(doc, "Capture el nombre", "Registre el nombre completo como debe aparecer en directorio, recibos y dispersión.")
    add_step(doc, "Elija el puesto", "Seleccione un puesto del catálogo controlado en Configuración.")
    add_step(doc, "Defina la vigencia", "Indique la fecha de alta. La persona sólo afecta periodos que coinciden con esa fecha.")
    add_step(doc, "Asigne sucursales", "Elija una, varias o todas como centros de costo. Una sucursal recibe 100 por ciento; varias se distribuyen equitativamente.")
    add_step(doc, "Configure el pago", "Registre salario mensual si corresponde y seleccione la nómina que paga el sueldo y la nómina que paga la comisión.")
    add_step(doc, "Capture el banco", "Registre el banco y una CLABE de exactamente 18 dígitos. El sistema rechaza longitudes menores o mayores.")
    add_step(doc, "Asigne el rol", "Elija el rol inicial y revise los módulos autorizados.")
    add_step(doc, "Guarde", "El empleado aparece en Roles y accesos y recibe una contraseña temporal para su primer acceso.")
    add_heading(doc, "Reglas de alta y baja", 2)
    add_bullet(doc, "Una alta nueva no modifica nóminas anteriores, salvo que la fecha de alta se coloque dentro de un periodo anterior.")
    add_bullet(doc, "La baja es lógica. Seleccione la fecha efectiva y no borre el registro.")
    add_bullet(doc, "La nómina del periodo de baja conserva los días trabajados y prorratea el sueldo cuando corresponde.")
    add_bullet(doc, "El salario mensual se divide en dos periodos: del 1 al 15 y del 16 al último día del mes.")
    add_bullet(doc, "Cambiar a una gerencia de área no elimina su historial; cambiarla a Ventas actualiza su portal y permisos desde la nueva vigencia.")
    add_heading(doc, "Edición y consulta", 2)
    add_body(
        doc,
        "Use el buscador para localizar personal por nombre, puesto, sucursal o rol. Cada renglón muestra puesto, tipo de nómina, acceso, banco, cuenta enmascarada, estado y vigencia. Antes de cambiar sueldo, puesto, sucursal o rol, confirme desde qué fecha debe aplicarse el cambio.",
    )

    doc.add_page_break()
    add_heading(doc, "5 Catálogos y estructura de nómina", 1)
    add_heading(doc, "Sucursales", 2)
    add_body(
        doc,
        "Configuración y Sucursales concentra altas, cambios, reactivaciones y bajas lógicas. Una sucursal activa aparece en empleados, centros de costo, movimientos, nóminas y reportes. La baja impide nuevas asignaciones pero conserva ventas, costos, recibos y gerencias históricas.",
    )
    add_body(
        doc,
        "La integración con el punto de venta está diseñada para altas y cambios automáticos mediante una clave externa única. En el prototipo, la sincronización y la persistencia son simuladas.",
    )
    add_heading(doc, "Puestos", 2)
    add_body(
        doc,
        "Registre cada puesto una sola vez con un nombre oficial. Los formularios de empleados y módulos sólo permiten elegir puestos de este catálogo. Esto unifica recibos, reportes y gastos por puesto.",
    )
    add_heading(doc, "Módulos de nómina", 2)
    add_body(
        doc,
        "Los módulos base se complementan con módulos personalizados. Para crear uno, escriba el nombre, seleccione por lo menos un concepto y asigne uno o varios puestos. El módulo aparece en Nómina, genera su periodo y reutiliza filtros, reportes, recibos y exportaciones.",
    )
    add_table(
        doc,
        ["Concepto configurable", "Uso"],
        [
            ["Sueldo", "Pago fijo o prorrateado según vigencia"],
            ["Comisión", "Pago variable según esquema y base con o sin IVA"],
            ["Bonos y multas", "Percepciones o deducciones autorizadas"],
            ["Ajustes", "Correcciones de más o de menos"],
            ["Préstamos y adelantos", "Descuentos o pagos asignados a un periodo"],
            ["Viáticos", "Importes autorizados y dirigidos a una nómina"],
        ],
        widths=[2.2, 4.4],
    )
    add_heading(doc, "Roles y permisos", 2)
    add_body(
        doc,
        "El rol forma parte del perfil de acceso y puede relacionarse con el puesto. El usuario máster no puede perder ningún módulo. Para otros roles, active únicamente los módulos y acciones necesarios. La navegación oculta lo no autorizado y la ruta directa también se rechaza.",
    )

    add_heading(doc, "6 Módulos de nómina", 1)
    add_table(
        doc,
        ["Módulo", "Qué incluye", "Regla principal"],
        [
            ["Consolidado", "Todas las nóminas, movimientos, cargas y centros de costo", "Los totales deben cuadrar con sucursales, puestos y tipos de nómina"],
            ["Salario fijo", "Sólo sueldo registrado y prorrateo", "No suma comisión ni movimientos variables"],
            ["Especialistas", "Sueldo y conceptos del módulo de especialistas", "Respeta vigencia, periodo y sucursal"],
            ["Comisiones", "Comisiones, bonos, deducciones y ajustes", "Nunca suma sueldo base, aun en puestos mixtos"],
            ["Comisión de kiosco", "Comisión mensual por sucursal y gerencia", "Mantiene récord histórico por sucursal"],
            ["Honorarios", "Subtotal, IVA, retenciones y neto", "Genera formato independiente"],
            ["Dispersión", "Personal de corridas cerradas listo para pago", "Es informativo y no realiza transferencias"],
            ["Personalizados", "Conceptos y puestos elegidos por el máster", "Alimentan consolidado y reportes una sola vez"],
        ],
        widths=[1.3, 2.65, 2.65],
        font_size=8.8,
    )
    add_heading(doc, "Puestos mixtos", 2)
    add_body(
        doc,
        "Un empleado puede tener sueldo base y comisión. El sueldo se dirige a su nómina salarial y la comisión a su nómina variable. El recibo de comisiones puede mostrar el sueldo como referencia, pero no lo suma al pago de ese módulo.",
    )
    add_heading(doc, "Comisión de kiosco", 2)
    add_body(
        doc,
        "El cálculo usa las ventas históricas de las sucursales incluidas en el esquema. Al final del módulo se muestra el récord más alto por sucursal y se actualiza con cada periodo agregado. Si no hay gerente, el historial conserva la leyenda Sin gerente. Al asignar una nueva gerencia, los periodos posteriores usan el nuevo nombre sin alterar los anteriores.",
    )

    doc.add_page_break()
    add_heading(doc, "7 Cálculo y cierre de una nómina", 1)
    add_heading(doc, "Preparar la corrida", 2)
    add_step(doc, "Abra el módulo", "Seleccione la nómina que desea calcular.")
    add_step(doc, "Revise el periodo", "Elija quincena o mes. Los reportes analíticos pueden ampliar el alcance a trimestre o año.")
    add_step(doc, "Elija la base", "En comisiones seleccione Con IVA o Sin IVA. La tabla siempre conserva ambas ventas para consulta.")
    add_step(doc, "Resuelva la distribución", "Si una persona vendió en varias sucursales, elija reparto parejo o proporcional a su participación de ventas.")
    add_step(doc, "Calcule", "Genere la corrida y revise cada empleado, cuenta, venta, comisión, deducción, ajuste, nómina, ISR, costo social y costo total.")
    add_step(doc, "Concilie", "Compare el total general con los subtotales por puesto, tipo de nómina y punto de venta.")
    add_step(doc, "Cierre", "Autorice la nómina cuando no existan diferencias. La corrida se protege y queda disponible para dispersión.")
    add_heading(doc, "Estados de la corrida", 2)
    add_table(
        doc,
        ["Estado", "Significado", "Acciones permitidas"],
        [
            ["Borrador", "Cálculo en preparación", "Editar configuración y volver a calcular"],
            ["Autorizada", "Nómina cerrada para pago", "Consultar, exportar, recibir aprobaciones y dispersar"],
            ["Pagada", "Proceso de pago marcado como terminado", "Consulta histórica y reportes"],
        ],
        widths=[1.25, 2.35, 3.0],
    )
    add_heading(doc, "Reapertura", 2)
    add_body(
        doc,
        "Una nómina cerrada no se puede modificar. Para reabrirla, el usuario autorizado debe ingresar el código máster. La corrida vuelve a Borrador, sale de Dispersión y debe conciliarse y autorizarse de nuevo.",
    )

    add_heading(doc, "8 Costo social e ISR", 1)
    add_body(
        doc,
        "Las cargas fiscales se administran por periodo. El interruptor del consolidado y la sección Aplicación global del periodo controlan de manera independiente el costo social y el ISR en nóminas, recibos, dispersión, consolidado y reportes.",
    )
    add_heading(doc, "Activar o apagar por periodo", 2)
    add_step(doc, "Seleccione el periodo", "Confirme la fecha inicial y final que desea modificar.")
    add_step(doc, "Cambie el interruptor", "Apague Costo social o ISR para convertir esa carga en cero en todos los módulos del periodo.")
    add_step(doc, "Verifique el resumen", "Cargas sociales y Costo general deben actualizarse de inmediato.")
    add_step(doc, "Cambie de periodo", "Confirme que otros meses o quincenas conservan su propia configuración.")
    add_heading(doc, "Configuración mixta por empleado", 2)
    add_body(
        doc,
        "En Configuración, abra Costo social e ISR. Seleccione el tipo de nómina. Use Copiar a todos para sincronizar una regla general o desactívelo para configurar cada empleado. Cada carga acepta Porcentaje o Monto fijo y una casilla para incluir o excluir a la persona.",
    )
    add_body(
        doc,
        "La regla global apagada prevalece sobre cualquier porcentaje o monto individual. Una corrida cerrada bloquea cambios fiscales del mismo periodo.",
    )

    doc.add_page_break()
    add_heading(doc, "9 Movimientos de nómina", 1)
    add_body(
        doc,
        "Movimientos registra percepciones y deducciones que no provienen directamente de la comisión o el sueldo. Cada movimiento conserva fecha de registro, empleado, tipo, monto, nómina, periodo, sucursal, concepto, destinos y estado.",
    )
    add_table(
        doc,
        ["Tipo", "Efecto habitual"],
        [
            ["Ajuste de más", "Suma al pago"],
            ["Ajuste de menos", "Resta del pago"],
            ["Bono", "Suma cuando está aprobado"],
            ["Multa", "Resta cuando está aprobada"],
            ["Préstamo", "Registra el adeudo y sus cuotas"],
            ["Pago de préstamo", "Reduce el saldo pendiente"],
            ["Sueldo base", "Sólo debe dirigirse al módulo salarial correspondiente"],
        ],
        widths=[2.0, 4.6],
    )
    add_heading(doc, "Crear y aprobar", 2)
    add_step(doc, "Pulse Nuevo movimiento", "Seleccione tipo, empleado y monto.")
    add_step(doc, "Seleccione el destino", "Indique la nómina, la corrida donde se pagará o descontará y la sucursal responsable.")
    add_step(doc, "Confirme los reportes", "Marque los destinos informativos que deberán reflejar el movimiento.")
    add_step(doc, "Envíe a revisión", "El borrador pasa a pendiente.")
    add_step(doc, "Apruebe", "Use la acción individual o seleccione varios pendientes y pulse Aprobar seleccionados.")
    add_heading(doc, "Selección masiva", 2)
    add_body(
        doc,
        "La casilla Seleccionar todos abarca los movimientos pendientes que cumplen los filtros, aunque estén en otra página. Antes de aprobar, revise la cantidad seleccionada. Una aprobación masiva actualiza nóminas, sucursales, recibos y reportes.",
    )
    add_heading(doc, "Movimientos de otro periodo", 2)
    add_body(
        doc,
        "El periodo activo está bloqueado por seguridad. Para capturar o mover un registro a otra quincena, ingrese el código máster, elija el nuevo destino y confirme que la fecha pertenece al corte. El movimiento vuelve a Borrador para una nueva aprobación.",
    )

    doc.add_page_break()
    add_heading(doc, "10 Préstamos y adelantos", 1)
    add_body(
        doc,
        "Cada solicitud debe indicar la nómina y el periodo donde comenzará el pago o descuento. Las cuotas posteriores conservan ese destino hasta liquidarse.",
    )
    add_table(
        doc,
        ["Regla inicial", "Valor predeterminado", "Resultado si no se cumple"],
        [
            ["Adelanto máximo", "50 por ciento de la comisión acumulada al día", "El sistema muestra una alerta y no permite continuar"],
            ["Adelantos por mes", "Máximo 2 solicitudes activas", "Se bloquea una solicitud adicional"],
            ["Cuotas por préstamo", "Máximo 6 cuotas", "No se acepta una cuota mayor"],
            ["Préstamos por trimestre", "Máximo 2 solicitudes activas", "Se bloquea una solicitud adicional"],
        ],
        widths=[1.7, 2.7, 2.2],
        font_size=9,
    )
    add_body(
        doc,
        "El usuario máster puede reducir o modificar estos límites desde Periodos y conceptos. Las solicitudes rechazadas no consumen el límite. Antes de autorizar, revise comisión disponible, saldo, cuotas, trimestre y nómina destino.",
    )

    add_heading(doc, "11 Recibos y autorizaciones del personal", 1)
    add_heading(doc, "Portal del empleado", 2)
    add_body(
        doc,
        "El empleado sólo ve su propia información. El recibo vigente muestra ventas, escala, comisión, bonos, multas, préstamos, ajustes, cuenta, pago y las cargas que correspondan. Los botones permiten Aprobar o Solicitar aclaración.",
    )
    add_step(doc, "Revise el corte vigente", "El portal no permite elegir ni aprobar otro periodo.")
    add_step(doc, "Compare los conceptos", "Valide ventas, base con o sin IVA, escala, descuentos, cuenta y total.")
    add_step(doc, "Decida", "Pulse Aprobar si todo es correcto o Solicitar aclaración si existe una diferencia.")
    add_step(doc, "Consulte el historial", "Después de aprobar, el recibo pasa al historial anual de sólo lectura. El portal muestra únicamente el año en curso.")
    add_heading(doc, "Vista del usuario máster", 2)
    add_body(
        doc,
        "Recibos muestra una columna de aprobación. Cuando el personal autoriza, aparece una palomita y el estado Aprobado por usuario. Si todos aprobaron, la conciliación se muestra en verde. Si falta alguien, aparece una alerta roja con los nombres pendientes.",
    )
    add_heading(doc, "Recibos gerenciales", 2)
    add_body(
        doc,
        "Una persona que vende y además comisiona como gerente conserva dos documentos separados. Su recibo personal de ventas permanece en Mi nómina. Cuando cierra el mes, Recibos gerenciales publica un recibo adicional con la comisión de kiosco, botones de aprobación y aclaración, y un historial anual independiente.",
    )

    add_heading(doc, "12 Reportes y conciliación", 1)
    add_heading(doc, "Consolidado", 2)
    add_body(
        doc,
        "El consolidado reúne todas las nóminas del periodo y muestra ventas con IVA, ventas sin IVA, sueldo, comisión, bonos, deducciones, ajustes, nómina, costo social, ISR, aprobación y costo total. La parte final concilia por puesto, tipo de nómina y punto de venta.",
    )
    add_status_line(doc, "Comparación exitosa", "El costo general coincide con la suma de todas las nóminas y no existen movimientos sin sucursal.", GREEN)
    add_status_line(doc, "Alerta de diferencia", "Existe un importe no asignado, un movimiento fuera de la sucursal correcta o una diferencia entre reportes. La alerta identifica el origen que debe revisarse.", RED)
    add_heading(doc, "Desglose por sucursal", 2)
    add_body(
        doc,
        "Incluye salario fijo, especialistas, comisiones, kiosco, honorarios, módulos personalizados y movimientos. Puede consultarse por mes, trimestre o año. Muestra ventas, nómina, costo social, ISR, costo integral, costo sobre venta, promedio por empleado, tendencia mensual y ranking de sucursales.",
    )
    add_heading(doc, "Gastos por puesto", 2)
    add_body(
        doc,
        "Filtre por fecha, nombre y puesto. Pulse el nombre del puesto para desplegar el detalle por empleado. El reporte muestra participación del costo, gráficas y el cambio porcentual frente al periodo equivalente del mes anterior.",
    )
    add_heading(doc, "Regla de conciliación", 2)
    add_table(
        doc,
        ["Dato", "Debe coincidir entre"],
        [
            ["Ventas con IVA", "Módulo de origen, consolidado y desglose por sucursal"],
            ["Ventas sin IVA", "Venta bruta dividida entre 1.16, cálculo y exportaciones"],
            ["Comisiones", "Cálculo por empleado, nómina de comisiones, recibo y reportes"],
            ["Nómina", "Suma de módulos sin duplicar sueldo o comisión"],
            ["Costo social e ISR", "Regla del periodo, empleados incluidos y todos los reportes"],
            ["Costo total", "Nómina más costo social más ISR"],
        ],
        widths=[1.7, 4.9],
    )

    add_heading(doc, "13 Dispersión y exportaciones", 1)
    add_heading(doc, "Dispersión de nómina", 2)
    add_body(
        doc,
        "Dispersión se habilita únicamente para corridas autorizadas o pagadas. Es un formato informativo previo a la transferencia y genera un documento separado por tipo de nómina, incluida la comisión de kiosco.",
    )
    add_table(
        doc,
        ["Campo", "Contenido"],
        [
            ["Identificación", "Apellido paterno, apellido materno y nombre"],
            ["Laboral", "Puesto y tipo de nómina"],
            ["Bancario", "Banco y CLABE interbancaria completa"],
            ["Pago", "Monto de pago, ISR, costo social y total"],
            ["Control", "Periodo, fecha de pago, estado y totales del formato"],
        ],
        widths=[1.55, 5.05],
    )
    add_heading(doc, "Imprimir y descargar", 2)
    add_step(doc, "Aplique los filtros", "Elija periodo, módulo, sucursal o empleado antes de exportar.")
    add_step(doc, "Revise los totales", "Confirme que la vista coincide con la corrida cerrada.")
    add_step(doc, "Elija el formato", "Use Imprimir, PDF o Excel. Cada salida debe conservar las columnas, los totales y la alineación de importes.")
    add_step(doc, "Proteja el archivo", "Guarde la exportación en una ubicación restringida por contener datos bancarios y de nómina.")
    add_body(
        doc,
        "PDF sirve para revisión, firma o archivo. Excel sirve para conciliación y preparación operativa. Imprimir abre la vista de impresión del navegador. Ninguna opción ejecuta una transferencia bancaria.",
    )

    add_heading(doc, "14 Procesos especiales", 1)
    add_heading(doc, "Distribución de un vendedor entre sucursales", 2)
    add_body(
        doc,
        "Antes de crear la nómina de comisiones, el sistema alerta cuando una persona vendió en dos o más sucursales. Reparto parejo asigna el mismo porcentaje a cada punto. Participación en ventas asigna más costo a la sucursal con mayor venta. La elección no cambia el pago del empleado.",
    )
    add_heading(doc, "Cambio de gerente", 2)
    add_body(
        doc,
        "Cierre la vigencia de la gerencia anterior y registre la fecha efectiva de la nueva asignación. El historial conserva el nombre previo. Si no existe gerente, continúe registrando el periodo como Sin gerente. Los permisos y el portal cambian desde la nueva asignación.",
    )
    add_heading(doc, "Venta con IVA y sin IVA", 2)
    add_body(
        doc,
        "La tabla muestra las dos cifras. Venta sin IVA equivale a la venta bruta dividida entre 1.16. Cuando la nómina usa Sin IVA, esa cifra se convierte en la base de la escala y de la comisión; cuando usa Con IVA, permanece como referencia.",
    )
    add_heading(doc, "Viáticos", 2)
    add_body(
        doc,
        "El máster configura conceptos y habilita qué empleados pueden solicitar. Cada solicitud indica nómina destino y sucursal. Sólo los viáticos aprobados afectan el pago y los reportes.",
    )

    add_heading(doc, "15 Controles antes del pago", 1)
    checklist = [
        ["1", "Periodo", "La fecha inicial y final son correctas en todos los módulos"],
        ["2", "Personal", "Altas, bajas, días laborados, puestos y sucursales están vigentes"],
        ["3", "Cuenta", "Cada CLABE contiene 18 dígitos y corresponde al empleado"],
        ["4", "Ventas", "Ventas con IVA y sin IVA coinciden con el origen"],
        ["5", "Comisiones", "Escala, base y porcentaje son correctos; no se sumó sueldo base"],
        ["6", "Movimientos", "Todos están aprobados y tienen nómina, periodo y sucursal"],
        ["7", "Préstamos", "Cuotas, saldo y límites están correctos"],
        ["8", "Cargas", "Costo social e ISR reflejan el interruptor y la selección individual"],
        ["9", "Conciliación", "Consolidado, sucursales, puestos y tipos de nómina cuadran"],
        ["10", "Aprobaciones", "El estatus identifica quién aprobó y quién está pendiente"],
        ["11", "Cierre", "La corrida está autorizada y protegida"],
        ["12", "Dispersión", "Los totales y las CLABE coinciden antes de exportar"],
    ]
    add_table(doc, ["Paso", "Control", "Validación"], checklist, widths=[0.65, 1.45, 4.5], font_size=8.8)

    doc.add_page_break()
    add_heading(doc, "16 Solución de problemas", 1)
    add_table(
        doc,
        ["Situación", "Qué revisar", "Acción"],
        [
            ["El sistema regresó al acceso", "Más de 3 minutos sin actividad", "Inicie sesión y complete el código secundario"],
            ["No aparece un módulo", "Permisos del rol", "El máster debe habilitar el módulo en Roles y accesos"],
            ["Un empleado no aparece", "Fecha de alta, baja, puesto y módulo", "Corrija la vigencia o la asignación desde Empleados"],
            ["La nómina no cuadra", "Movimientos, distribución y filtros", "Abra la alerta de conciliación y revise el origen señalado"],
            ["ISR o costo social no aparece", "Interruptor del periodo y casilla individual", "Active la carga correcta en el periodo seleccionado"],
            ["No se puede editar", "Corrida autorizada o pagada", "Use reapertura con código máster y vuelva a conciliar"],
            ["No aparece en dispersión", "Estado de la corrida", "Autorice o marque como pagada la nómina"],
            ["La CLABE es rechazada", "Cantidad y tipo de caracteres", "Capture exactamente 18 dígitos, sin espacios"],
            ["No se puede pedir un adelanto", "Comisión acumulada y solicitudes del mes", "Reduzca el monto o revise la política configurada"],
            ["Un cambio desapareció", "Recarga o restauración de la demo", "Repita la prueba; la versión actual no persiste datos"],
        ],
        widths=[1.75, 2.0, 2.85],
        font_size=8.6,
    )

    add_heading(doc, "17 Glosario y rutas del sistema", 1)
    add_heading(doc, "Glosario", 2)
    add_table(
        doc,
        ["Término", "Definición"],
        [
            ["Corrida", "Cálculo de una nómina para un módulo y periodo específicos"],
            ["Nómina base", "Pago antes de agregar costo social e ISR"],
            ["Costo social", "Carga patronal configurada como porcentaje o monto fijo"],
            ["ISR", "Provisión fiscal configurada como porcentaje o monto fijo"],
            ["Costo total", "Nómina base más costo social más ISR"],
            ["Centro de costo", "Sucursal que absorbe todo o parte del costo laboral"],
            ["Dispersión", "Listado bancario informativo de una nómina cerrada"],
            ["Vigencia", "Fecha desde la cual una asignación o regla se aplica"],
            ["Aclaración", "Solicitud del empleado cuando no acepta un concepto de su recibo"],
        ],
        widths=[1.6, 5.0],
    )
    add_heading(doc, "Rutas principales", 2)
    routes = [
        ["Inicio", "/", "Consolidado general"],
        ["Empleados", "/empleados", "Directorio y alta"],
        ["Salario fijo", "/nomina-salario-fijo", "Nómina salarial"],
        ["Especialistas", "/nomina-especialistas", "Nómina de especialistas"],
        ["Comisiones", "/nomina-comisiones", "Nómina variable"],
        ["Kiosco", "/nomina-comision-kiosco", "Comisión gerencial por sucursal"],
        ["Honorarios", "/nomina-honorarios", "Pago por honorarios"],
        ["Dispersión", "/dispersion-nomina", "Formatos para pago"],
        ["Cálculo", "/calculo-comisiones", "Detalle auditable de comisiones"],
        ["Movimientos", "/movimientos", "Percepciones y deducciones"],
        ["Préstamos", "/prestamos-adelantos", "Solicitudes y cuotas"],
        ["Configuración", "/configuracion", "Periodos, conceptos y políticas"],
        ["Módulos", "/modulos-nomina", "Nóminas personalizadas"],
        ["Sucursales", "/sucursales", "Catálogo de puntos de venta"],
        ["Puestos", "/puestos", "Catálogo de puestos"],
        ["Esquemas", "/esquemas", "Escalas de comisión"],
        ["Esquemas sucursal", "/esquemas-sucursal", "Comisión de kiosco y gerencias"],
        ["Bonos y multas", "/bonos-multas", "Conceptos y autorización"],
        ["Viáticos", "/viaticos", "Configuración y comprobación"],
        ["Accesos", "/accesos", "Roles, permisos y códigos"],
        ["Gastos por puesto", "/reportes/gastos-por-puesto", "Análisis por puesto"],
        ["Desglose sucursal", "/reportes/desglose-sucursal", "Análisis por punto de venta"],
        ["Recibos", "/recibos", "Control general de recibos"],
        ["Recibos gerenciales", "/recibos-kiosco", "Comisión gerencial"],
        ["Mi nómina", "/mi-nomina", "Portal privado del empleado"],
    ]
    add_table(doc, ["Módulo", "Ruta", "Uso"], routes, widths=[1.65, 2.35, 2.6], font_size=8.3)

    add_heading(doc, "Cierre del proceso", 2)
    add_body(
        doc,
        "El proceso termina cuando la corrida está conciliada, autorizada, revisada por el personal y disponible en Dispersión. Conserve el PDF final y el Excel de control en una ubicación restringida, junto con las aclaraciones y la evidencia de cualquier reapertura.",
    )

    doc.core_properties.title = "Manual de uso del sistema de nómina"
    doc.core_properties.subject = "Operación del portal de nómina de Keysar Cosmetics"
    doc.core_properties.author = "Keysar Cosmetics"
    doc.core_properties.keywords = "nómina, manual, empleados, comisiones, reportes, dispersión"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_manual()
